#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, asdict
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SOURCE_CODE = 'CBI_RIAL_FEE_1405_PROVISIONAL'
SOURCE_TYPE_CODE = 'CIRCULAR'
ISSUER_CODE = 'CBI'
CIRCULAR_NO = 'PROVISIONAL-RIAL-FEE-1405'
POLICY_CODE = 'CBI_RIAL_BANKING_1405_PROVISIONAL'
POLICY_VERSION_NO = '1405.P1'
CLASSIFICATION_CODE = 'CBI_1405_RIAL_PROVISIONAL'
CREATED_BY = 'CBI_1405_PROVISIONAL_IMPORT'
EFFECTIVE_FROM = '2026-09-10'
OLD_SOURCE_CODE = 'CBI_FEE_1404_04_35500'
OLD_CLASSIFICATION = 'CBI_1404'
OLD_VERSION_NO = '1.0'
OLD_ARCHIVE_TO = '2026-09-09'

SECTION_MAP = {
    '1': ('GAR', 'CBI_RIAL_1405_GUARANTEE_GROUP', 'ضمانتنامه‌ها'),
    '2': ('REM', 'CBI_RIAL_1405_REMITTANCE_GROUP', 'حواله‌ها'),
    '3': ('SDB', 'CBI_RIAL_1405_SAFE_DEPOSIT_GROUP', 'صندوق‌های اجاره‌ای'),
    '4': ('SEC', 'CBI_RIAL_1405_SECURITIES_GROUP', 'مدیریت اوراق بهادار مشتریان'),
    '5': ('COL', 'CBI_RIAL_1405_COLLECTION_GROUP', 'وصول بروات'),
    '6': ('ACC', 'CBI_RIAL_1405_ACCOUNT_GROUP', 'انواع حساب‌ها'),
    '7': ('APP', 'CBI_RIAL_1405_APPRAISAL_GROUP', 'ارزیابی اموال منقول و غیرمنقول'),
    '8': ('CRD', 'CBI_RIAL_1405_CREDIT_GROUP', 'اعتبارات'),
    '9': ('OTH', 'CBI_RIAL_1405_OTHER_GROUP', 'سایر خدمات'),
}

OLD_RIAL_FEATURES = [
    'CBI_GUARANTEE_FEE_GROUP',
    'CBI_REMITTANCE_FEE_GROUP',
    'CBI_SAFE_DEPOSIT_FEE_GROUP',
    'CBI_SECURITIES_FEE_GROUP',
    'CBI_BILL_COLLECTION_FEE_GROUP',
    'CBI_CURRENT_ACCOUNT_FEE_GROUP',
    'CBI_SAVINGS_DEPOSIT_FEE_GROUP',
    'CBI_CERTIFICATE_FEE_GROUP',
    'CBI_STATEMENT_FEE_GROUP',
    'CBI_ACCOUNT_SERVICE_FEE_GROUP',
    'CBI_APPRAISAL_FEE_GROUP',
    'CBI_CREDIT_FEE_GROUP',
    'CBI_OTHER_SERVICE_FEE_GROUP',
]
OLD_ELECTRONIC_FEATURES = [
    'CBI_ELECTRONIC_SERVICE_FEE_GROUP',
    'CBI_ELECTRONIC_LC_RIAL_FEE_GROUP',
    'CBI_ELECTRONIC_GUARANTEE_RIAL_FEE_GROUP',
    'CBI_ELECTRONIC_BILL_FEE_GROUP',
]


def clean(value: Any) -> str | None:
    if value is None:
        return None
    s = unicodedata.normalize('NFKC', str(value))
    s = s.translate(str.maketrans({'ي':'ی','ى':'ی','ك':'ک','ة':'ه','ۀ':'ه','\u200c':' ','\u200d':' ','\u200e':' ','\u200f':' ','\ufeff':' ','\xa0':' ','ـ':''}))
    s = s.replace('\r',' ').replace('\n',' ').replace('\t',' ')
    s = re.sub(r'\s+', ' ', s).strip()
    return s or None


def dec(v: Any) -> Decimal | None:
    if v is None or v == '':
        return None
    return Decimal(str(v))


def sql_str(v: str | None) -> str:
    return 'NULL' if v is None else "'" + v.replace("'", "''") + "'"


def sql_num(v: Decimal | int | float | None) -> str:
    if v is None:
        return 'NULL'
    d = Decimal(str(v))
    s = format(d, 'f')
    if '.' in s:
        s = s.rstrip('0').rstrip('.')
    return s or '0'


def slug(code: str) -> str:
    return re.sub(r'[^0-9A-Za-z]+', '_', code).strip('_').upper()


def normalized_rate(percent: Any, permille: Any) -> Decimal | None:
    if percent not in (None, ''):
        return dec(percent) / Decimal('100')
    if permille not in (None, ''):
        return dec(permille) / Decimal('1000')
    return None


def component_numeric(c: dict[str, Any]) -> Decimal | None:
    t = c['rule_type']
    if t == 'FIXED_AMOUNT':
        return c['amount']
    if t == 'RATE':
        return normalized_rate(c['percent'], c['permille'])
    if t == 'NO_FEE':
        return Decimal('0')
    if t == 'RATE_UNIT':
        return c['amount']
    return None


def component_node_type(c: dict[str, Any]) -> str:
    if c['rule_type'] in {'FIXED_AMOUNT','RATE','NO_FEE','RATE_UNIT','DEPOSIT_MULTIPLIER'}:
        return 'CONSTANT'
    return 'EXTERNAL_VALUE'


def choose_strategy(r: dict[str, Any], comps: list[dict[str, Any]]) -> tuple[str,str,Decimal|None,Decimal|None,Decimal|None,Decimal|None,str|None]:
    if len(comps) != 1 or r['composite'] == 'بله' or '/' in (r['rule_type'] or ''):
        return ('COMPOSITE','Other',None,None,None,None,None)
    c = comps[0]
    t = c['rule_type']
    if t == 'FIXED_AMOUNT':
        if c['unit']:
            return ('PER_UNIT','PerUnit',c['amount'],None,c['min_fee'],c['max_fee'],None)
        return ('FIXED','Flat',c['amount'],None,c['min_fee'],c['max_fee'],None)
    if t == 'NO_FEE':
        return ('FIXED','Flat',Decimal('0'),None,None,None,None)
    if t == 'RATE':
        rate = normalized_rate(c['percent'],c['permille'])
        annual = c['frequency'] == 'سالانه'
        if annual:
            strategy = 'ANNUALIZED_PERCENTAGE'
        elif c['min_fee'] is not None and c['max_fee'] is not None:
            strategy = 'PERCENTAGE_FLOOR_CAP'
        elif c['min_fee'] is not None:
            strategy = 'PERCENTAGE_WITH_FLOOR'
        elif c['max_fee'] is not None:
            strategy = 'PERCENTAGE_WITH_CAP'
        else:
            strategy = 'PERCENTAGE'
        if c['min_fee'] is not None and c['max_fee'] is not None:
            basis='Other' if annual else 'Other'
        elif c['min_fee'] is not None:
            basis='RateWithMinimumAmount'
        elif c['max_fee'] is not None:
            basis='RateWithMaximumAmount'
        else:
            basis='Percentage'
        return (strategy,basis,None,rate,c['min_fee'],c['max_fee'],'YEAR' if annual else None)
    return ('EXTERNAL_VALUE','Other',None,None,c['min_fee'],c['max_fee'],'YEAR' if c['frequency']=='سالانه' else None)


@dataclass
class Tariff:
    tariff_code: str
    section_no: str
    section_code: str
    feature_code: str
    section_name: str
    service_name: str
    fee_text: str
    fee_code: str
    category_code: str
    strategy: str
    basis_type: str
    fixed_amount: Decimal|None
    rate_value: Decimal|None
    min_fee: Decimal|None
    max_fee: Decimal|None
    rate_period_code: str|None
    component_count: int
    composite: str
    page: int
    source_row: int
    config_hash: str


def load_source(xlsx: Path) -> tuple[list[Tariff], list[dict[str,Any]], dict[str, Any]]:
    wb=load_workbook(xlsx,data_only=True,read_only=True)
    expected={'کارمزدها','کارمزدها_ساختاری','اجزای_محاسبه','راهنمای_فیلدها','یادداشت‌ها و تبصره‌ها','راهنما'}
    if set(wb.sheetnames) != expected:
        raise ValueError(f'Unexpected sheets: {wb.sheetnames}')

    ws=wb['کارمزدها_ساختاری']; headers=[c.value for c in ws[1]]; ix={h:i for i,h in enumerate(headers)}
    raw_rows=[]
    for row in ws.iter_rows(min_row=2,values_only=True):
        raw_rows.append({h:row[i] for h,i in ix.items()})
    if len(raw_rows)!=152:
        raise ValueError(f'Expected 152 structured tariffs, got {len(raw_rows)}')

    wc=wb['اجزای_محاسبه']; hc=[c.value for c in wc[1]]; ci={h:i for i,h in enumerate(hc)}
    components=[]
    for row in wc.iter_rows(min_row=2,values_only=True):
        d={h:row[i] for h,i in ci.items()}
        components.append({
            'component_id': clean(d['شناسه جزء']),
            'source_row': int(d['ردیف منبع Excel']),
            'page': int(d['صفحه PDF']),
            'section_no': clean(d['شماره سرفصل']),
            'tariff_code': clean(d['ردیف تعرفه']),
            'service_name': clean(d['نوع خدمت']),
            'sequence': int(d['ترتیب جزء']),
            'rule_type': clean(d['نوع قاعده']),
            'status': clean(d['وضعیت']),
            'customer_group': clean(d['گروه مشتری']),
            'percent': dec(d['درصد (%)']),
            'permille': dec(d['در هزار (‰)']),
            'rate_limit': clean(d['محدودیت نرخ']),
            'amount': dec(d['مبلغ (ریال)']),
            'min_fee': dec(d['حداقل مبلغ (ریال)']),
            'max_fee': dec(d['حداکثر مبلغ (ریال)']),
            'lower_bound': dec(d['حد پایین بازه (ریال)']),
            'upper_bound': dec(d['حد بالای بازه (ریال)']),
            'range_basis': clean(d['مبنای بازه']),
            'basis': clean(d['مبنای محاسبه']),
            'unit': clean(d['واحد اخذ']),
            'frequency': clean(d['تناوب']),
            'extra_cost': clean(d['هزینه اضافی']),
            'reference': clean(d['مرجع تعرفه/قاعده']),
            'condition': clean(d['شرط اجرا']),
            'description': clean(d['شرح جزء/قاعده']),
            'original_fee_text': clean(d['متن اصلی کارمزد مصوب']),
            'source_ref': clean(d['مرجع منبع']),
        })
    if len(components)!=180:
        raise ValueError(f'Expected 180 source component rows, got {len(components)}')

    by_tariff=defaultdict(list)
    for c in components:
        by_tariff[c['tariff_code']].append(c)
    for v in by_tariff.values(): v.sort(key=lambda x:x['sequence'])

    tariffs=[]
    seen=set()
    for r in raw_rows:
        code=clean(r['ردیف تعرفه'])
        if not code or code in seen:
            raise ValueError(f'Missing/duplicate tariff code: {code}')
        seen.add(code)
        sec=clean(r['شماره سرفصل'])
        if sec not in SECTION_MAP: raise ValueError(f'Unknown section: {sec}')
        sec_code,feature_code,sec_short=SECTION_MAP[sec]
        comps=by_tariff[code]
        if len(comps)!=int(r['تعداد اجزای محاسبه'] or 0):
            raise ValueError(f'Component count mismatch for {code}: {len(comps)} != {r["تعداد اجزای محاسبه"]}')
        strategy,basis,fixed,rate,min_fee,max_fee,period=choose_strategy({
            'composite':clean(r['قاعده مرکب؟']), 'rule_type':clean(r['نوع قاعده'])
        },comps)
        fee_code=f'CBI1405R_{sec_code}_{slug(code)}'
        payload={
            'tariff_code':code,'section':clean(r['سرفصل خدمات بانکی']),'service':clean(r['نوع خدمت']),
            'fee_text':clean(r['متن اصلی کارمزد مصوب']),'components':[
                {k:(str(v) if isinstance(v,Decimal) else v) for k,v in c.items()} for c in comps
            ],
        }
        ch=hashlib.sha256(json.dumps(payload,ensure_ascii=False,sort_keys=True).encode('utf-8')).hexdigest()
        tariffs.append(Tariff(
            tariff_code=code, section_no=sec, section_code=sec_code, feature_code=feature_code,
            section_name=clean(r['سرفصل خدمات بانکی']) or sec_short,
            service_name=clean(r['نوع خدمت']) or code,
            fee_text=clean(r['متن اصلی کارمزد مصوب']) or '',
            fee_code=fee_code, category_code=f'CBI1405R_{sec_code}',
            strategy=strategy,basis_type=basis,fixed_amount=fixed,rate_value=rate,min_fee=min_fee,max_fee=max_fee,
            rate_period_code=period,component_count=len(comps),composite=clean(r['قاعده مرکب؟']) or 'خیر',
            page=int(r['صفحه PDF']),source_row=int(r['ردیف منبع Excel']),config_hash=ch
        ))
    meta={'rule_strategies':dict(Counter(t.strategy for t in tariffs)),
          'sections':dict(Counter(t.section_no for t in tariffs)),
          'source_components':len(components),
          'tariffs':len(tariffs)}
    return tariffs,components,meta


def component_description(c: dict[str,Any]) -> str:
    parts=[f"جزء منبع {c['component_id']}", f"نوع={c['rule_type']}"]
    if c['status']: parts.append(f"وضعیت={c['status']}")
    if c['customer_group']: parts.append(f"گروه مشتری={c['customer_group']}")
    if c['rate_limit']: parts.append(f"محدودیت نرخ={c['rate_limit']}")
    if c['basis']: parts.append(f"مبنا={c['basis']}")
    if c['unit']: parts.append(f"واحد={c['unit']}")
    if c['frequency']: parts.append(f"تناوب={c['frequency']}")
    if c['lower_bound'] is not None or c['upper_bound'] is not None:
        parts.append(f"بازه={c['lower_bound'] or ''}..{c['upper_bound'] or ''}")
    if c['range_basis']: parts.append(f"مبنای بازه={c['range_basis']}")
    if c['extra_cost']: parts.append(f"هزینه اضافی={c['extra_cost']}")
    if c['reference']: parts.append(f"مرجع={c['reference']}")
    if c['condition']: parts.append(f"شرط={c['condition']}")
    if c['description']: parts.append(c['description'])
    return ' | '.join(parts)[:990]


def write_csv(path: Path, rows: list[dict[str,Any]], fields: list[str]):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader()
        for row in rows:
            w.writerow({k:(str(row.get(k)) if isinstance(row.get(k),Decimal) else row.get(k)) for k in fields})


def generate_sql(tariffs:list[Tariff], components:list[dict[str,Any]], xlsx_hash:str, pdf_hash:str) -> str:
    by_tariff=defaultdict(list)
    for c in components: by_tariff[c['tariff_code']].append(c)
    lines=[]
    a=lines.append
    a('-- ============================================================================')
    a('-- CBI Rial Banking Fees 1405 - PROVISIONAL versioned import')
    a(f'-- Workbook SHA256: {xlsx_hash}')
    a(f'-- PDF SHA256     : {pdf_hash}')
    a('-- Official circular number/date were not provided. This dataset is intentionally marked PROVISIONAL.')
    a('-- Replaces only the non-electronic rial-banking subset of CBI 1404; electronic-fee definitions remain untouched.')
    a('-- ============================================================================')
    a('SET DEFINE OFF;')
    a('WHENEVER SQLERROR EXIT SQL.SQLCODE ROLLBACK;')
    a('ALTER SESSION SET CURRENT_SCHEMA = FEE;')
    a('')
    a('PROMPT [CBI1405-P] Validating prerequisite CBI 1404 import ...')
    a("DECLARE v NUMBER; BEGIN SELECT COUNT(*) INTO v FROM FEE_REGULATORY_SOURCE WHERE SOURCE_CODE='CBI_FEE_1404_04_35500'; IF v<>1 THEN RAISE_APPLICATION_ERROR(-20201,'CBI 1404 source missing/duplicated'); END IF; END;")
    a('/')
    a('')
    # source
    a('PROMPT [CBI1405-P] Upserting provisional regulatory source ...')
    doc_ref=("پیوست (کارمزدهای ریالی).pdf | کارمزدهای ریالی تفکیک ساختاری برای درج دیتابیس v3 2026-09-10.xlsx | "
             f"XLSX_SHA256={xlsx_hash}")[:490]
    a('MERGE INTO FEE_REGULATORY_SOURCE t')
    a('USING (SELECT '+', '.join([
        f"{sql_str(SOURCE_CODE)} SOURCE_CODE",f"{sql_str(SOURCE_TYPE_CODE)} SOURCE_TYPE_CODE",f"{sql_str(ISSUER_CODE)} ISSUER_CODE",
        f"{sql_str(CIRCULAR_NO)} CIRCULAR_NO",sql_str('کارمزد خدمات بانکی ریالی - مصوبه جدید (مرجع موقت)')+' TITLE_FA',
        sql_str('CBI Rial Banking Fees - Provisional Reference')+' TITLE_EN',f"DATE '{EFFECTIVE_FROM}' ISSUE_DATE",f"DATE '{EFFECTIVE_FROM}' EFFECTIVE_FROM",
        f"{sql_str(doc_ref)} DOCUMENT_REF",f"{sql_str('SHA256:'+pdf_hash)} DOCUMENT_HASH",sql_str('PROVISIONAL')+' STATUS_CODE'])+' FROM dual) s')
    a('ON (t.SOURCE_CODE=s.SOURCE_CODE)')
    a("WHEN MATCHED THEN UPDATE SET t.SOURCE_TYPE_CODE=s.SOURCE_TYPE_CODE,t.ISSUER_CODE=s.ISSUER_CODE,t.CIRCULAR_NO=s.CIRCULAR_NO,t.TITLE_FA=s.TITLE_FA,t.TITLE_EN=s.TITLE_EN,t.ISSUE_DATE=s.ISSUE_DATE,t.EFFECTIVE_FROM=s.EFFECTIVE_FROM,t.EFFECTIVE_TO=NULL,t.DOCUMENT_REF=s.DOCUMENT_REF,t.DOCUMENT_HASH=s.DOCUMENT_HASH,t.STATUS_CODE=s.STATUS_CODE,t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
    a('WHEN NOT MATCHED THEN INSERT (SOURCE_CODE,SOURCE_TYPE_CODE,ISSUER_CODE,CIRCULAR_NO,TITLE_FA,TITLE_EN,ISSUE_DATE,EFFECTIVE_FROM,DOCUMENT_REF,DOCUMENT_HASH,STATUS_CODE,CREATED_BY)')
    a("VALUES (s.SOURCE_CODE,s.SOURCE_TYPE_CODE,s.ISSUER_CODE,s.CIRCULAR_NO,s.TITLE_FA,s.TITLE_EN,s.ISSUE_DATE,s.EFFECTIVE_FROM,s.DOCUMENT_REF,s.DOCUMENT_HASH,s.STATUS_CODE,'CBI_1405_PROVISIONAL_IMPORT');")
    a('')
    # policy
    a('PROMPT [CBI1405-P] Upserting provisional policy and version ...')
    a('MERGE INTO FEE_POLICY_SET t USING (SELECT '+sql_str(POLICY_CODE)+" POLICY_CODE,"+sql_str('سیاست کارمزد خدمات بانکی ریالی ۱۴۰۵ - موقت')+" NAME_FA,"+sql_str('CBI Rial Banking Fee Policy 1405 - Provisional')+" NAME_EN,'REGULATORY' POLICY_TYPE_CODE,'CBI' OWNER_ORG_CODE FROM dual) s ON (t.POLICY_CODE=s.POLICY_CODE)")
    a("WHEN MATCHED THEN UPDATE SET t.NAME_FA=s.NAME_FA,t.NAME_EN=s.NAME_EN,t.POLICY_TYPE_CODE=s.POLICY_TYPE_CODE,t.DESCRIPTION='مرجع موقت تا زمان دریافت نامه رسمی؛ جایگزین فقط دامنه کارمزدهای خدمات بانکی ریالی غیرالکترونیکی',t.OWNER_ORG_CODE=s.OWNER_ORG_CODE,t.STATUS_CODE='ACTIVE',t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
    a("WHEN NOT MATCHED THEN INSERT (POLICY_CODE,NAME_FA,NAME_EN,POLICY_TYPE_CODE,DESCRIPTION,OWNER_ORG_CODE,STATUS_CODE,CREATED_BY) VALUES (s.POLICY_CODE,s.NAME_FA,s.NAME_EN,s.POLICY_TYPE_CODE,'مرجع موقت تا زمان دریافت نامه رسمی؛ جایگزین فقط دامنه کارمزدهای خدمات بانکی ریالی غیرالکترونیکی',s.OWNER_ORG_CODE,'ACTIVE','CBI_1405_PROVISIONAL_IMPORT');")
    a('MERGE INTO FEE_POLICY_VERSION t USING (SELECT (SELECT POLICY_SET_ID FROM FEE_POLICY_SET WHERE POLICY_CODE='+sql_str(POLICY_CODE)+') POLICY_SET_ID,'+sql_str(POLICY_VERSION_NO)+" VERSION_NO FROM dual) s ON (t.POLICY_SET_ID=s.POLICY_SET_ID AND t.VERSION_NO=s.VERSION_NO)")
    a(f"WHEN MATCHED THEN UPDATE SET t.VERSION_TYPE_CODE='MAJOR',t.STATUS_CODE='ACTIVE',t.EFFECTIVE_FROM=DATE '{EFFECTIVE_FROM}',t.EFFECTIVE_TO=NULL,t.APPROVED_AT=SYSTIMESTAMP,t.APPROVED_BY='PROVISIONAL',t.ACTIVATED_AT=SYSTIMESTAMP,t.ACTIVATED_BY='PROVISIONAL',t.POLICY_HASH={sql_str(xlsx_hash)},t.CHANGE_SUMMARY='مصوبه موقت برگرفته از پیوست ۸ صفحه‌ای و فایل ساختاری ۱۵۲ تعرفه؛ تاریخ/شماره رسمی در دسترس نیست',t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
    a(f"WHEN NOT MATCHED THEN INSERT (POLICY_SET_ID,VERSION_NO,VERSION_TYPE_CODE,STATUS_CODE,EFFECTIVE_FROM,APPROVED_AT,APPROVED_BY,ACTIVATED_AT,ACTIVATED_BY,POLICY_HASH,CHANGE_SUMMARY,CREATED_BY) VALUES (s.POLICY_SET_ID,s.VERSION_NO,'MAJOR','ACTIVE',DATE '{EFFECTIVE_FROM}',SYSTIMESTAMP,'PROVISIONAL',SYSTIMESTAMP,'PROVISIONAL',{sql_str(xlsx_hash)},'مصوبه موقت برگرفته از پیوست ۸ صفحه‌ای و فایل ساختاری ۱۵۲ تعرفه؛ تاریخ/شماره رسمی در دسترس نیست','CBI_1405_PROVISIONAL_IMPORT');")
    a('')
    # archive old subset only
    feats=', '.join(sql_str(x) for x in OLD_RIAL_FEATURES)
    a('PROMPT [CBI1405-P] Closing only prior non-electronic rial fee definition versions ...')
    a("UPDATE FEE_DEFINITION_VERSION dv SET dv.STATUS_CODE='SUPERSEDED', dv.EFFECTIVE_TO=DATE '"+OLD_ARCHIVE_TO+"', dv.UPDATED_AT=SYSTIMESTAMP, dv.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT', dv.RECORD_VERSION=dv.RECORD_VERSION+1 WHERE dv.REGULATORY_SOURCE_ID=(SELECT REGULATORY_SOURCE_ID FROM FEE_REGULATORY_SOURCE WHERE SOURCE_CODE='"+OLD_SOURCE_CODE+"') AND dv.VERSION_NO='"+OLD_VERSION_NO+"' AND EXISTS (SELECT 1 FROM FEE_DEFINITION d JOIN FEE_FEATURE f ON f.FEE_FEATURE_ID=d.FEE_FEATURE_ID WHERE d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID AND f.FEATURE_CODE IN ("+feats+") ) AND (dv.EFFECTIVE_TO IS NULL OR dv.EFFECTIVE_TO>DATE '"+OLD_ARCHIVE_TO+"');")
    a("UPDATE FEE_CALCULATION_RULE r SET r.EFFECTIVE_TO=DATE '"+OLD_ARCHIVE_TO+"', r.UPDATED_AT=SYSTIMESTAMP, r.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT', r.RECORD_VERSION=r.RECORD_VERSION+1 WHERE EXISTS (SELECT 1 FROM FEE_DEFINITION_VERSION dv JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID JOIN FEE_FEATURE f ON f.FEE_FEATURE_ID=d.FEE_FEATURE_ID JOIN FEE_REGULATORY_SOURCE rs ON rs.REGULATORY_SOURCE_ID=dv.REGULATORY_SOURCE_ID WHERE dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID AND rs.SOURCE_CODE='"+OLD_SOURCE_CODE+"' AND dv.VERSION_NO='"+OLD_VERSION_NO+"' AND f.FEATURE_CODE IN ("+feats+") ) AND (r.EFFECTIVE_TO IS NULL OR r.EFFECTIVE_TO>DATE '"+OLD_ARCHIVE_TO+"');")
    a('')
    # features
    a('PROMPT [CBI1405-P] Upserting 9 source section features ...')
    for sec,(sec_code,feature_code,short) in SECTION_MAP.items():
        section_name=next(t.section_name for t in tariffs if t.section_no==sec)
        a(f"MERGE INTO FEE_FEATURE t USING (SELECT {sql_str(feature_code)} FEATURE_CODE,{sql_str(section_name)} NAME_FA,'Other' FEE_TYPE_CODE,{sql_str('CBI1405R_'+sec_code)} DEFAULT_CATEGORY_CODE FROM dual) s ON (t.FEATURE_CODE=s.FEATURE_CODE)")
        a("WHEN MATCHED THEN UPDATE SET t.NAME_FA=s.NAME_FA,t.FEE_TYPE_CODE=s.FEE_TYPE_CODE,t.DEFAULT_CATEGORY_CODE=s.DEFAULT_CATEGORY_CODE,t.DESCRIPTION='گروه تعرفه مصوبه موقت کارمزد خدمات بانکی ریالی ۱۴۰۵',t.IS_ACTIVE='Y',t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
        a("WHEN NOT MATCHED THEN INSERT (FEATURE_CODE,NAME_FA,FEE_TYPE_CODE,DEFAULT_CATEGORY_CODE,DESCRIPTION,IS_ACTIVE,CREATED_BY) VALUES (s.FEATURE_CODE,s.NAME_FA,s.FEE_TYPE_CODE,s.DEFAULT_CATEGORY_CODE,'گروه تعرفه مصوبه موقت کارمزد خدمات بانکی ریالی ۱۴۰۵','Y','CBI_1405_PROVISIONAL_IMPORT');")
    a('')
    # definitions
    a(f'PROMPT [CBI1405-P] Upserting {len(tariffs)} fee definitions, versions and calculation rules ...')
    for t in tariffs:
        desc=(f"منبع موقت کارمزد خدمات بانکی ریالی ۱۴۰۵ | سرفصل {t.section_no}: {t.section_name} | ردیف تعرفه: {t.tariff_code} | صفحه PDF: {t.page} | متن مصوب: {t.fee_text}")[:1980]
        a(f"MERGE INTO FEE_DEFINITION t USING (SELECT (SELECT FEE_FEATURE_ID FROM FEE_FEATURE WHERE FEATURE_CODE={sql_str(t.feature_code)}) FEE_FEATURE_ID,{sql_str(t.fee_code)} FEE_CODE,{sql_str(t.service_name)} NAME_FA,{sql_str(t.category_code)} CATEGORY_CODE,{sql_str(CLASSIFICATION_CODE)} CLASSIFICATION_CODE,{sql_str(desc)} DESCRIPTION FROM dual) s ON (t.FEE_CODE=s.FEE_CODE)")
        a("WHEN MATCHED THEN UPDATE SET t.FEE_FEATURE_ID=s.FEE_FEATURE_ID,t.NAME_FA=s.NAME_FA,t.CATEGORY_CODE=s.CATEGORY_CODE,t.CLASSIFICATION_CODE=s.CLASSIFICATION_CODE,t.DESCRIPTION=s.DESCRIPTION,t.IS_ACTIVE='Y',t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
        a("WHEN NOT MATCHED THEN INSERT (FEE_FEATURE_ID,FEE_CODE,NAME_FA,CATEGORY_CODE,CLASSIFICATION_CODE,DESCRIPTION,IS_ACTIVE,CREATED_BY) VALUES (s.FEE_FEATURE_ID,s.FEE_CODE,s.NAME_FA,s.CATEGORY_CODE,s.CLASSIFICATION_CODE,s.DESCRIPTION,'Y','CBI_1405_PROVISIONAL_IMPORT');")
        reason=(f"متن کارمزد مصوب: {t.fee_text} | مرجع موقت؛ شماره و تاریخ رسمی بعداً جایگزین شود")[:1980]
        a(f"MERGE INTO FEE_DEFINITION_VERSION t USING (SELECT (SELECT FEE_DEFINITION_ID FROM FEE_DEFINITION WHERE FEE_CODE={sql_str(t.fee_code)}) FEE_DEFINITION_ID,(SELECT POLICY_VERSION_ID FROM FEE_POLICY_VERSION pv JOIN FEE_POLICY_SET ps ON ps.POLICY_SET_ID=pv.POLICY_SET_ID WHERE ps.POLICY_CODE={sql_str(POLICY_CODE)} AND pv.VERSION_NO={sql_str(POLICY_VERSION_NO)}) POLICY_VERSION_ID,(SELECT REGULATORY_SOURCE_ID FROM FEE_REGULATORY_SOURCE WHERE SOURCE_CODE={sql_str(SOURCE_CODE)}) REGULATORY_SOURCE_ID FROM dual) s ON (t.FEE_DEFINITION_ID=s.FEE_DEFINITION_ID AND t.VERSION_NO={sql_str(POLICY_VERSION_NO)})")
        a(f"WHEN MATCHED THEN UPDATE SET t.POLICY_VERSION_ID=s.POLICY_VERSION_ID,t.REGULATORY_SOURCE_ID=s.REGULATORY_SOURCE_ID,t.REGULATORY_TARIFF_CODE={sql_str(t.tariff_code)},t.VERSION_TYPE_CODE='MAJOR',t.STATUS_CODE='ACTIVE',t.FEE_PLAN_NAME={sql_str('تعرفه ریالی ۱۴۰۵ - ردیف '+t.tariff_code)},t.FEE_PLAN_TYPE_CODE='REGULATORY',t.FEE_REASON={sql_str(reason)},t.EFFECTIVE_FROM=DATE '{EFFECTIVE_FROM}',t.EFFECTIVE_TO=NULL,t.APPROVED_AT=SYSTIMESTAMP,t.APPROVED_BY='PROVISIONAL',t.ACTIVATED_AT=SYSTIMESTAMP,t.ACTIVATED_BY='PROVISIONAL',t.CONFIG_HASH={sql_str(t.config_hash)},t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
        a(f"WHEN NOT MATCHED THEN INSERT (FEE_DEFINITION_ID,POLICY_VERSION_ID,REGULATORY_SOURCE_ID,REGULATORY_TARIFF_CODE,VERSION_NO,VERSION_TYPE_CODE,STATUS_CODE,FEE_PLAN_NAME,FEE_PLAN_TYPE_CODE,FEE_REASON,EFFECTIVE_FROM,APPROVED_AT,APPROVED_BY,ACTIVATED_AT,ACTIVATED_BY,CONFIG_HASH,CREATED_BY) VALUES (s.FEE_DEFINITION_ID,s.POLICY_VERSION_ID,s.REGULATORY_SOURCE_ID,{sql_str(t.tariff_code)},{sql_str(POLICY_VERSION_NO)},'MAJOR','ACTIVE',{sql_str('تعرفه ریالی ۱۴۰۵ - ردیف '+t.tariff_code)},'REGULATORY',{sql_str(reason)},DATE '{EFFECTIVE_FROM}',SYSTIMESTAMP,'PROVISIONAL',SYSTIMESTAMP,'PROVISIONAL',{sql_str(t.config_hash)},'CBI_1405_PROVISIONAL_IMPORT');")
        rule_code='TARIFF_'+slug(t.tariff_code)
        calc_desc=(f"قاعده ساختاری از فایل پیوست؛ ردیف {t.tariff_code}; تعداد اجزا {t.component_count}; متن: {t.fee_text}")[:1980]
        a(f"MERGE INTO FEE_CALCULATION_RULE t USING (SELECT (SELECT v.FEE_DEFINITION_VERSION_ID FROM FEE_DEFINITION_VERSION v JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=v.FEE_DEFINITION_ID WHERE d.FEE_CODE={sql_str(t.fee_code)} AND v.VERSION_NO={sql_str(POLICY_VERSION_NO)}) FEE_DEFINITION_VERSION_ID FROM dual) s ON (t.FEE_DEFINITION_VERSION_ID=s.FEE_DEFINITION_VERSION_ID AND t.RULE_CODE={sql_str(rule_code)})")
        a(f"WHEN MATCHED THEN UPDATE SET t.NAME_FA={sql_str(t.service_name)},t.PRIORITY_NO=100,t.CALCULATION_STRATEGY_CODE={sql_str(t.strategy)},t.BASIS_TYPE_CODE={sql_str(t.basis_type)},t.FIXED_AMOUNT={sql_num(t.fixed_amount)},t.RATE_VALUE={sql_num(t.rate_value)},t.MIN_FEE_AMOUNT={sql_num(t.min_fee)},t.MAX_FEE_AMOUNT={sql_num(t.max_fee)},t.RATE_PERIOD_CODE={sql_str(t.rate_period_code)},t.CURRENCY_CODE='IRR',t.ROUNDING_MODE_CODE='HALF_UP',t.ROUNDING_SCALE=0,t.EFFECTIVE_FROM=DATE '{EFFECTIVE_FROM}',t.EFFECTIVE_TO=NULL,t.IS_ACTIVE='Y',t.DESCRIPTION={sql_str(calc_desc)},t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
        a(f"WHEN NOT MATCHED THEN INSERT (FEE_DEFINITION_VERSION_ID,RULE_CODE,NAME_FA,PRIORITY_NO,CALCULATION_STRATEGY_CODE,BASIS_TYPE_CODE,FIXED_AMOUNT,RATE_VALUE,MIN_FEE_AMOUNT,MAX_FEE_AMOUNT,RATE_PERIOD_CODE,CURRENCY_CODE,ROUNDING_MODE_CODE,ROUNDING_SCALE,EFFECTIVE_FROM,IS_ACTIVE,DESCRIPTION,CREATED_BY) VALUES (s.FEE_DEFINITION_VERSION_ID,{sql_str(rule_code)},{sql_str(t.service_name)},100,{sql_str(t.strategy)},{sql_str(t.basis_type)},{sql_num(t.fixed_amount)},{sql_num(t.rate_value)},{sql_num(t.min_fee)},{sql_num(t.max_fee)},{sql_str(t.rate_period_code)},'IRR','HALF_UP',0,DATE '{EFFECTIVE_FROM}','Y',{sql_str(calc_desc)},'CBI_1405_PROVISIONAL_IMPORT');")
        # Generic input for directly executable rate/per-unit rules.
        if t.strategy in {'PERCENTAGE','PERCENTAGE_WITH_FLOOR','PERCENTAGE_WITH_CAP','PERCENTAGE_FLOOR_CAP','ANNUALIZED_PERCENTAGE'}:
            basis_name=by_tariff[t.tariff_code][0]['basis'] or 'مبلغ مبنای محاسبه'
            a(f"MERGE INTO FEE_INPUT_DEFINITION t USING (SELECT (SELECT CALCULATION_RULE_ID FROM FEE_CALCULATION_RULE r WHERE r.FEE_DEFINITION_VERSION_ID=(SELECT v.FEE_DEFINITION_VERSION_ID FROM FEE_DEFINITION_VERSION v JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=v.FEE_DEFINITION_ID WHERE d.FEE_CODE={sql_str(t.fee_code)} AND v.VERSION_NO={sql_str(POLICY_VERSION_NO)}) AND r.RULE_CODE={sql_str(rule_code)}) CALCULATION_RULE_ID FROM dual) s ON (t.CALCULATION_RULE_ID=s.CALCULATION_RULE_ID AND t.INPUT_CODE='BASE_AMOUNT')")
            a(f"WHEN MATCHED THEN UPDATE SET t.NAME_FA={sql_str(basis_name)},t.DATA_TYPE_CODE='NUMBER',t.UNIT_CODE='IRR',t.MANDATORY_FLAG='Y',t.DISPLAY_ORDER=1,t.DESCRIPTION='مبنای محاسبه استخراج‌شده از فایل ساختاری',t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
            a(f"WHEN NOT MATCHED THEN INSERT (CALCULATION_RULE_ID,INPUT_CODE,NAME_FA,DATA_TYPE_CODE,UNIT_CODE,MANDATORY_FLAG,DISPLAY_ORDER,DESCRIPTION,CREATED_BY) VALUES (s.CALCULATION_RULE_ID,'BASE_AMOUNT',{sql_str(basis_name)},'NUMBER','IRR','Y',1,'مبنای محاسبه استخراج‌شده از فایل ساختاری','CBI_1405_PROVISIONAL_IMPORT');")
        # Tiered appraisal rows: preserve the three source ranges in FEE_CALCULATION_TIER.
        ranged=[c for c in by_tariff[t.tariff_code] if c['lower_bound'] is not None or c['upper_bound'] is not None]
        if len(ranged)==3 and all(c['range_basis']=='مبلغ ارزیابی' for c in ranged):
            for input_code,input_name,order in [('APPRAISAL_AMOUNT','مبلغ ارزیابی',1),('OFFICIAL_EXPERT_TARIFF','تعرفه کارشناس رسمی دادگستری',2)]:
                a(f"MERGE INTO FEE_INPUT_DEFINITION t USING (SELECT (SELECT CALCULATION_RULE_ID FROM FEE_CALCULATION_RULE r WHERE r.FEE_DEFINITION_VERSION_ID=(SELECT v.FEE_DEFINITION_VERSION_ID FROM FEE_DEFINITION_VERSION v JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=v.FEE_DEFINITION_ID WHERE d.FEE_CODE={sql_str(t.fee_code)} AND v.VERSION_NO={sql_str(POLICY_VERSION_NO)}) AND r.RULE_CODE={sql_str(rule_code)}) CALCULATION_RULE_ID FROM dual) s ON (t.CALCULATION_RULE_ID=s.CALCULATION_RULE_ID AND t.INPUT_CODE={sql_str(input_code)})")
                a(f"WHEN MATCHED THEN UPDATE SET t.NAME_FA={sql_str(input_name)},t.DATA_TYPE_CODE='NUMBER',t.UNIT_CODE='IRR',t.MANDATORY_FLAG='Y',t.DISPLAY_ORDER={order},t.DESCRIPTION='ورودی صریح برای قاعده ارزیابی پلکانی مصوبه موقت ۱۴۰۵',t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
                a(f"WHEN NOT MATCHED THEN INSERT (CALCULATION_RULE_ID,INPUT_CODE,NAME_FA,DATA_TYPE_CODE,UNIT_CODE,MANDATORY_FLAG,DISPLAY_ORDER,DESCRIPTION,CREATED_BY) VALUES (s.CALCULATION_RULE_ID,{sql_str(input_code)},{sql_str(input_name)},'NUMBER','IRR','Y',{order},'ورودی صریح برای قاعده ارزیابی پلکانی مصوبه موقت ۱۴۰۵','CBI_1405_PROVISIONAL_IMPORT');")
            for tier_no,c in enumerate(ranged,1):
                if c['rule_type']=='FIXED_AMOUNT':
                    tier_strategy='FIXED'; fixed=c['amount']; rate=None; basis_code='WHOLE_AMOUNT'
                elif c['rule_type']=='RATE':
                    tier_strategy='COMPOSITE' if c['reference'] else 'PERCENTAGE'; fixed=None; rate=normalized_rate(c['percent'],c['permille']); basis_code='EXCESS_OVER_LOWER_BOUND'
                else:
                    tier_strategy='EXTERNAL_VALUE'; fixed=None; rate=None; basis_code='EXCESS_OVER_LOWER_BOUND'
                tier_name=c['description'] or f"بازه {tier_no}"
                a(f"MERGE INTO FEE_CALCULATION_TIER t USING (SELECT (SELECT CALCULATION_RULE_ID FROM FEE_CALCULATION_RULE r WHERE r.FEE_DEFINITION_VERSION_ID=(SELECT v.FEE_DEFINITION_VERSION_ID FROM FEE_DEFINITION_VERSION v JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=v.FEE_DEFINITION_ID WHERE d.FEE_CODE={sql_str(t.fee_code)} AND v.VERSION_NO={sql_str(POLICY_VERSION_NO)}) AND r.RULE_CODE={sql_str(rule_code)}) CALCULATION_RULE_ID FROM dual) s ON (t.CALCULATION_RULE_ID=s.CALCULATION_RULE_ID AND t.TIER_NO={tier_no})")
                a(f"WHEN MATCHED THEN UPDATE SET t.TIER_NAME_FA={sql_str(tier_name[:190])},t.LOWER_BOUND={sql_num(c['lower_bound'])},t.UPPER_BOUND={sql_num(c['upper_bound'])},t.BOUND_UNIT_CODE='IRR',t.TIER_BASIS_CODE={sql_str(basis_code)},t.TIER_STRATEGY_CODE={sql_str(tier_strategy)},t.FIXED_AMOUNT={sql_num(fixed)},t.RATE_VALUE={sql_num(rate)},t.MIN_FEE_AMOUNT={sql_num(c['min_fee'])},t.MAX_FEE_AMOUNT={sql_num(c['max_fee'])},t.EFFECTIVE_FROM=DATE '{EFFECTIVE_FROM}',t.EFFECTIVE_TO=NULL,t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
                a(f"WHEN NOT MATCHED THEN INSERT (CALCULATION_RULE_ID,TIER_NO,TIER_NAME_FA,LOWER_BOUND,UPPER_BOUND,BOUND_UNIT_CODE,TIER_BASIS_CODE,TIER_STRATEGY_CODE,FIXED_AMOUNT,RATE_VALUE,MIN_FEE_AMOUNT,MAX_FEE_AMOUNT,EFFECTIVE_FROM,CREATED_BY) VALUES (s.CALCULATION_RULE_ID,{tier_no},{sql_str(tier_name[:190])},{sql_num(c['lower_bound'])},{sql_num(c['upper_bound'])},'IRR',{sql_str(basis_code)},{sql_str(tier_strategy)},{sql_num(fixed)},{sql_num(rate)},{sql_num(c['min_fee'])},{sql_num(c['max_fee'])},DATE '{EFFECTIVE_FROM}','CBI_1405_PROVISIONAL_IMPORT');")
        elif t.strategy=='PER_UNIT':
            unit=by_tariff[t.tariff_code][0]['unit'] or 'واحد'
            a(f"MERGE INTO FEE_INPUT_DEFINITION t USING (SELECT (SELECT CALCULATION_RULE_ID FROM FEE_CALCULATION_RULE r WHERE r.FEE_DEFINITION_VERSION_ID=(SELECT v.FEE_DEFINITION_VERSION_ID FROM FEE_DEFINITION_VERSION v JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=v.FEE_DEFINITION_ID WHERE d.FEE_CODE={sql_str(t.fee_code)} AND v.VERSION_NO={sql_str(POLICY_VERSION_NO)}) AND r.RULE_CODE={sql_str(rule_code)}) CALCULATION_RULE_ID FROM dual) s ON (t.CALCULATION_RULE_ID=s.CALCULATION_RULE_ID AND t.INPUT_CODE='UNIT_COUNT')")
            a(f"WHEN MATCHED THEN UPDATE SET t.NAME_FA={sql_str('تعداد '+unit)},t.DATA_TYPE_CODE='NUMBER',t.UNIT_CODE='COUNT',t.MANDATORY_FLAG='Y',t.DISPLAY_ORDER=1,t.DESCRIPTION={sql_str('واحد منبع: '+unit)},t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
            a(f"WHEN NOT MATCHED THEN INSERT (CALCULATION_RULE_ID,INPUT_CODE,NAME_FA,DATA_TYPE_CODE,UNIT_CODE,MANDATORY_FLAG,DISPLAY_ORDER,DESCRIPTION,CREATED_BY) VALUES (s.CALCULATION_RULE_ID,'UNIT_COUNT',{sql_str('تعداد '+unit)},'NUMBER','COUNT','Y',1,{sql_str('واحد منبع: '+unit)},'CBI_1405_PROVISIONAL_IMPORT');")
        # Source components: flat auditable representation; not used to invent a formula tree.
        for c in by_tariff[t.tariff_code]:
            node=component_node_type(c); num=component_numeric(c)
            ref=f"SRC:{c['component_id']}:{c['rule_type']}"[:190]
            desc_c=component_description(c)
            a(f"MERGE INTO FEE_RULE_COMPONENT t USING (SELECT (SELECT CALCULATION_RULE_ID FROM FEE_CALCULATION_RULE r WHERE r.FEE_DEFINITION_VERSION_ID=(SELECT v.FEE_DEFINITION_VERSION_ID FROM FEE_DEFINITION_VERSION v JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=v.FEE_DEFINITION_ID WHERE d.FEE_CODE={sql_str(t.fee_code)} AND v.VERSION_NO={sql_str(POLICY_VERSION_NO)}) AND r.RULE_CODE={sql_str(rule_code)}) CALCULATION_RULE_ID FROM dual) s ON (t.CALCULATION_RULE_ID=s.CALCULATION_RULE_ID AND t.SEQUENCE_NO={c['sequence']})")
            a(f"WHEN MATCHED THEN UPDATE SET t.PARENT_RULE_COMPONENT_ID=NULL,t.NODE_TYPE_CODE={sql_str(node)},t.OPERATOR_CODE=NULL,t.INPUT_CODE=NULL,t.CONSTANT_NUMBER={sql_num(num)},t.CONSTANT_TEXT={sql_str(c['original_fee_text'][:490] if c['original_fee_text'] else None)},t.REFERENCE_CODE={sql_str(ref)},t.DESCRIPTION={sql_str(desc_c)},t.UPDATED_AT=SYSTIMESTAMP,t.UPDATED_BY='CBI_1405_PROVISIONAL_IMPORT',t.RECORD_VERSION=t.RECORD_VERSION+1")
            a(f"WHEN NOT MATCHED THEN INSERT (CALCULATION_RULE_ID,PARENT_RULE_COMPONENT_ID,SEQUENCE_NO,NODE_TYPE_CODE,OPERATOR_CODE,INPUT_CODE,CONSTANT_NUMBER,CONSTANT_TEXT,REFERENCE_CODE,DESCRIPTION,CREATED_BY) VALUES (s.CALCULATION_RULE_ID,NULL,{c['sequence']},{sql_str(node)},NULL,NULL,{sql_num(num)},{sql_str(c['original_fee_text'][:490] if c['original_fee_text'] else None)},{sql_str(ref)},{sql_str(desc_c)},'CBI_1405_PROVISIONAL_IMPORT');")
        a('')
    a('PROMPT [CBI1405-P] Import staged. Run verification before COMMIT.')
    return '\n'.join(lines)+'\n'



def generate_reconciliation(tariffs:list[Tariff], components:list[dict[str,Any]], xlsx_hash:str, pdf_hash:str) -> str:
    """Generate a source-to-Oracle acceptance verifier.

    Unlike the structural verifier, this embeds the normalized source contract and
    compares every tariff, every source component, every executable input and every
    generated tier against the rows stored in Oracle. No source file access is
    required at execution time.
    """
    by_tariff=defaultdict(list)
    for c in components:
        by_tariff[c['tariff_code']].append(c)
    for rows in by_tariff.values():
        rows.sort(key=lambda x:x['sequence'])

    lines=[]
    a=lines.append
    a('-- ============================================================================')
    a('-- CBI Rial Banking Fees 1405 - PROVISIONAL source-to-Oracle reconciliation')
    a(f'-- Workbook SHA256: {xlsx_hash}')
    a(f'-- PDF SHA256     : {pdf_hash}')
    a('-- Generated from the same normalized source contract as the FIX98 importer.')
    a('-- This script is read-only. It raises ORA-20260 if any business-value mismatch exists.')
    a('-- ============================================================================')
    a('SET DEFINE OFF;')
    a('WHENEVER SQLERROR EXIT SQL.SQLCODE ROLLBACK;')
    a('ALTER SESSION SET CURRENT_SCHEMA = FEE;')
    a('SET SERVEROUTPUT ON SIZE UNLIMITED;')
    a('')
    a('DECLARE')
    a('  v_errors NUMBER := 0;')
    a('  v_tariffs_checked NUMBER := 0;')
    a('  v_components_checked NUMBER := 0;')
    a('  v_inputs_checked NUMBER := 0;')
    a('  v_tiers_checked NUMBER := 0;')
    a('')
    a('  FUNCTION same_text(a VARCHAR2, b VARCHAR2) RETURN BOOLEAN IS')
    a('  BEGIN')
    a('    IF a IS NULL AND b IS NULL THEN RETURN TRUE; END IF;')
    a('    IF a IS NULL OR b IS NULL THEN RETURN FALSE; END IF;')
    a('    RETURN a = b;')
    a('  END;')
    a('')
    a('  FUNCTION same_num(a NUMBER, b NUMBER) RETURN BOOLEAN IS')
    a('  BEGIN')
    a('    IF a IS NULL AND b IS NULL THEN RETURN TRUE; END IF;')
    a('    IF a IS NULL OR b IS NULL THEN RETURN FALSE; END IF;')
    a('    RETURN a = b;')
    a('  END;')
    a('')
    a('  FUNCTION same_date(a DATE, b DATE) RETURN BOOLEAN IS')
    a('  BEGIN')
    a('    IF a IS NULL AND b IS NULL THEN RETURN TRUE; END IF;')
    a('    IF a IS NULL OR b IS NULL THEN RETURN FALSE; END IF;')
    a('    RETURN a = b;')
    a('  END;')
    a('')
    a('  FUNCTION show_text(v VARCHAR2) RETURN VARCHAR2 IS')
    a("  BEGIN RETURN CASE WHEN v IS NULL THEN '<NULL>' ELSE SUBSTR(v,1,900) END; END;")
    a('')
    a('  FUNCTION show_num(v NUMBER) RETURN VARCHAR2 IS')
    a("  BEGIN RETURN CASE WHEN v IS NULL THEN '<NULL>' ELSE TO_CHAR(v) END; END;")
    a('')
    a('  PROCEDURE fail(p_scope VARCHAR2, p_key VARCHAR2, p_field VARCHAR2, p_expected VARCHAR2, p_actual VARCHAR2) IS')
    a('  BEGIN')
    a('    v_errors := v_errors + 1;')
    a("    DBMS_OUTPUT.PUT_LINE('[MISMATCH] '||p_scope||' '||p_key||' '||p_field||' expected='||show_text(p_expected)||' actual='||show_text(p_actual));")
    a('  END;')
    a('')
    a('  PROCEDURE check_text(p_scope VARCHAR2, p_key VARCHAR2, p_field VARCHAR2, p_expected VARCHAR2, p_actual VARCHAR2) IS')
    a('  BEGIN IF NOT same_text(p_expected,p_actual) THEN fail(p_scope,p_key,p_field,p_expected,p_actual); END IF; END;')
    a('')
    a('  PROCEDURE check_num(p_scope VARCHAR2, p_key VARCHAR2, p_field VARCHAR2, p_expected NUMBER, p_actual NUMBER) IS')
    a('  BEGIN IF NOT same_num(p_expected,p_actual) THEN fail(p_scope,p_key,p_field,show_num(p_expected),show_num(p_actual)); END IF; END;')
    a('')
    a('  PROCEDURE check_date(p_scope VARCHAR2, p_key VARCHAR2, p_field VARCHAR2, p_expected DATE, p_actual DATE) IS')
    a("  BEGIN IF NOT same_date(p_expected,p_actual) THEN fail(p_scope,p_key,p_field,CASE WHEN p_expected IS NULL THEN '<NULL>' ELSE TO_CHAR(p_expected,'YYYY-MM-DD') END,CASE WHEN p_actual IS NULL THEN '<NULL>' ELSE TO_CHAR(p_actual,'YYYY-MM-DD') END); END IF; END;")
    a('')
    a('  PROCEDURE check_tariff(')
    a('    p_fee_code VARCHAR2, p_tariff_code VARCHAR2, p_rule_code VARCHAR2, p_name_fa VARCHAR2,')
    a('    p_feature_code VARCHAR2, p_category_code VARCHAR2, p_config_hash VARCHAR2,')
    a('    p_strategy VARCHAR2, p_basis VARCHAR2, p_fixed NUMBER, p_rate NUMBER,')
    a('    p_min NUMBER, p_max NUMBER, p_period VARCHAR2, p_component_count NUMBER')
    a('  ) IS')
    a('    a_name FEE.FEE_DEFINITION.NAME_FA%TYPE; a_feature FEE.FEE_FEATURE.FEATURE_CODE%TYPE; a_category FEE.FEE_DEFINITION.CATEGORY_CODE%TYPE; a_class FEE.FEE_DEFINITION.CLASSIFICATION_CODE%TYPE;')
    a('    a_tariff FEE.FEE_DEFINITION_VERSION.REGULATORY_TARIFF_CODE%TYPE; a_version FEE.FEE_DEFINITION_VERSION.VERSION_NO%TYPE; a_status FEE.FEE_DEFINITION_VERSION.STATUS_CODE%TYPE; a_hash FEE.FEE_DEFINITION_VERSION.CONFIG_HASH%TYPE;')
    a('    a_source FEE.FEE_REGULATORY_SOURCE.SOURCE_CODE%TYPE; a_policy FEE.FEE_POLICY_SET.POLICY_CODE%TYPE; a_rule FEE.FEE_CALCULATION_RULE.RULE_CODE%TYPE; a_strategy FEE.FEE_CALCULATION_RULE.CALCULATION_STRATEGY_CODE%TYPE; a_basis FEE.FEE_CALCULATION_RULE.BASIS_TYPE_CODE%TYPE;')
    a('    a_fixed NUMBER; a_rate NUMBER; a_min NUMBER; a_max NUMBER; a_period FEE.FEE_CALCULATION_RULE.RATE_PERIOD_CODE%TYPE;')
    a('    a_currency FEE.FEE_CALCULATION_RULE.CURRENCY_CODE%TYPE; a_active FEE.FEE_CALCULATION_RULE.IS_ACTIVE%TYPE; a_from DATE; a_to DATE; a_components NUMBER;')
    a('  BEGIN')
    a('    v_tariffs_checked := v_tariffs_checked + 1;')
    a('    BEGIN')
    a('      SELECT d.NAME_FA,f.FEATURE_CODE,d.CATEGORY_CODE,d.CLASSIFICATION_CODE,')
    a('             dv.REGULATORY_TARIFF_CODE,dv.VERSION_NO,dv.STATUS_CODE,dv.CONFIG_HASH,')
    a('             rs.SOURCE_CODE,ps.POLICY_CODE,r.RULE_CODE,r.CALCULATION_STRATEGY_CODE,r.BASIS_TYPE_CODE,')
    a('             r.FIXED_AMOUNT,r.RATE_VALUE,r.MIN_FEE_AMOUNT,r.MAX_FEE_AMOUNT,r.RATE_PERIOD_CODE,')
    a('             r.CURRENCY_CODE,r.IS_ACTIVE,r.EFFECTIVE_FROM,r.EFFECTIVE_TO,')
    a("             (SELECT COUNT(*) FROM FEE_RULE_COMPONENT c WHERE c.CALCULATION_RULE_ID=r.CALCULATION_RULE_ID AND c.REFERENCE_CODE LIKE 'SRC:%')")
    a('        INTO a_name,a_feature,a_category,a_class,a_tariff,a_version,a_status,a_hash,a_source,a_policy,a_rule,a_strategy,a_basis,')
    a('             a_fixed,a_rate,a_min,a_max,a_period,a_currency,a_active,a_from,a_to,a_components')
    a('        FROM FEE_DEFINITION d')
    a('        JOIN FEE_FEATURE f ON f.FEE_FEATURE_ID=d.FEE_FEATURE_ID')
    a('        JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_ID=d.FEE_DEFINITION_ID')
    a('        JOIN FEE_REGULATORY_SOURCE rs ON rs.REGULATORY_SOURCE_ID=dv.REGULATORY_SOURCE_ID')
    a('        JOIN FEE_POLICY_VERSION pv ON pv.POLICY_VERSION_ID=dv.POLICY_VERSION_ID')
    a('        JOIN FEE_POLICY_SET ps ON ps.POLICY_SET_ID=pv.POLICY_SET_ID')
    a('        JOIN FEE_CALCULATION_RULE r ON r.FEE_DEFINITION_VERSION_ID=dv.FEE_DEFINITION_VERSION_ID')
    a(f"       WHERE d.FEE_CODE=p_fee_code AND dv.VERSION_NO='{POLICY_VERSION_NO}' AND r.RULE_CODE=p_rule_code;")
    a('    EXCEPTION')
    a("      WHEN NO_DATA_FOUND THEN fail('TARIFF',p_fee_code,'ROW','present','missing'); RETURN;")
    a("      WHEN TOO_MANY_ROWS THEN fail('TARIFF',p_fee_code,'ROW','one row','multiple rows'); RETURN;")
    a('    END;')
    a("    check_text('TARIFF',p_fee_code,'NAME_FA',p_name_fa,a_name);")
    a("    check_text('TARIFF',p_fee_code,'FEATURE_CODE',p_feature_code,a_feature);")
    a("    check_text('TARIFF',p_fee_code,'CATEGORY_CODE',p_category_code,a_category);")
    a(f"    check_text('TARIFF',p_fee_code,'CLASSIFICATION_CODE','{CLASSIFICATION_CODE}',a_class);")
    a("    check_text('TARIFF',p_fee_code,'REGULATORY_TARIFF_CODE',p_tariff_code,a_tariff);")
    a(f"    check_text('TARIFF',p_fee_code,'VERSION_NO','{POLICY_VERSION_NO}',a_version);")
    a("    check_text('TARIFF',p_fee_code,'STATUS_CODE','ACTIVE',a_status);")
    a("    check_text('TARIFF',p_fee_code,'CONFIG_HASH',p_config_hash,a_hash);")
    a(f"    check_text('TARIFF',p_fee_code,'SOURCE_CODE','{SOURCE_CODE}',a_source);")
    a(f"    check_text('TARIFF',p_fee_code,'POLICY_CODE','{POLICY_CODE}',a_policy);")
    a("    check_text('TARIFF',p_fee_code,'RULE_CODE',p_rule_code,a_rule);")
    a("    check_text('TARIFF',p_fee_code,'STRATEGY',p_strategy,a_strategy);")
    a("    check_text('TARIFF',p_fee_code,'BASIS_TYPE',p_basis,a_basis);")
    a("    check_num('TARIFF',p_fee_code,'FIXED_AMOUNT',p_fixed,a_fixed);")
    a("    check_num('TARIFF',p_fee_code,'RATE_VALUE',p_rate,a_rate);")
    a("    check_num('TARIFF',p_fee_code,'MIN_FEE_AMOUNT',p_min,a_min);")
    a("    check_num('TARIFF',p_fee_code,'MAX_FEE_AMOUNT',p_max,a_max);")
    a("    check_text('TARIFF',p_fee_code,'RATE_PERIOD_CODE',p_period,a_period);")
    a("    check_text('TARIFF',p_fee_code,'CURRENCY_CODE','IRR',a_currency);")
    a("    check_text('TARIFF',p_fee_code,'IS_ACTIVE','Y',a_active);")
    a(f"    check_date('TARIFF',p_fee_code,'EFFECTIVE_FROM',DATE '{EFFECTIVE_FROM}',a_from);")
    a("    check_date('TARIFF',p_fee_code,'EFFECTIVE_TO',NULL,a_to);")
    a("    check_num('TARIFF',p_fee_code,'SOURCE_COMPONENT_COUNT',p_component_count,a_components);")
    a('  END;')
    a('')
    a('  PROCEDURE check_component(p_fee_code VARCHAR2, p_sequence NUMBER, p_node VARCHAR2, p_number NUMBER, p_text VARCHAR2, p_ref VARCHAR2, p_desc VARCHAR2) IS')
    a('    a_node FEE.FEE_RULE_COMPONENT.NODE_TYPE_CODE%TYPE; a_number NUMBER; a_text FEE.FEE_RULE_COMPONENT.CONSTANT_TEXT%TYPE; a_ref FEE.FEE_RULE_COMPONENT.REFERENCE_CODE%TYPE; a_desc FEE.FEE_RULE_COMPONENT.DESCRIPTION%TYPE;')
    a('    a_parent NUMBER; a_operator FEE.FEE_RULE_COMPONENT.OPERATOR_CODE%TYPE; a_input FEE.FEE_RULE_COMPONENT.INPUT_CODE%TYPE;')
    a("    k VARCHAR2(200) := p_fee_code||'#'||p_sequence;")
    a('  BEGIN')
    a('    v_components_checked := v_components_checked + 1;')
    a('    BEGIN')
    a('      SELECT c.NODE_TYPE_CODE,c.CONSTANT_NUMBER,c.CONSTANT_TEXT,c.REFERENCE_CODE,c.DESCRIPTION,c.PARENT_RULE_COMPONENT_ID,c.OPERATOR_CODE,c.INPUT_CODE')
    a('        INTO a_node,a_number,a_text,a_ref,a_desc,a_parent,a_operator,a_input')
    a('        FROM FEE_RULE_COMPONENT c')
    a('        JOIN FEE_CALCULATION_RULE r ON r.CALCULATION_RULE_ID=c.CALCULATION_RULE_ID')
    a('        JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID')
    a('        JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID')
    a(f"       WHERE d.FEE_CODE=p_fee_code AND dv.VERSION_NO='{POLICY_VERSION_NO}' AND c.SEQUENCE_NO=p_sequence;")
    a('    EXCEPTION')
    a("      WHEN NO_DATA_FOUND THEN fail('COMPONENT',k,'ROW','present','missing'); RETURN;")
    a("      WHEN TOO_MANY_ROWS THEN fail('COMPONENT',k,'ROW','one row','multiple rows'); RETURN;")
    a('    END;')
    a("    check_text('COMPONENT',k,'NODE_TYPE_CODE',p_node,a_node);")
    a("    check_num('COMPONENT',k,'CONSTANT_NUMBER',p_number,a_number);")
    a("    check_text('COMPONENT',k,'CONSTANT_TEXT',p_text,a_text);")
    a("    check_text('COMPONENT',k,'REFERENCE_CODE',p_ref,a_ref);")
    a("    check_text('COMPONENT',k,'DESCRIPTION',p_desc,a_desc);")
    a("    IF a_parent IS NOT NULL THEN fail('COMPONENT',k,'PARENT_RULE_COMPONENT_ID','<NULL>',show_num(a_parent)); END IF;")
    a("    check_text('COMPONENT',k,'OPERATOR_CODE',NULL,a_operator);")
    a("    check_text('COMPONENT',k,'INPUT_CODE',NULL,a_input);")
    a('  END;')
    a('')
    a('  PROCEDURE check_input(p_fee_code VARCHAR2, p_input_code VARCHAR2, p_name VARCHAR2, p_unit VARCHAR2, p_order NUMBER) IS')
    a('    a_name FEE.FEE_INPUT_DEFINITION.NAME_FA%TYPE; a_type FEE.FEE_INPUT_DEFINITION.DATA_TYPE_CODE%TYPE; a_unit FEE.FEE_INPUT_DEFINITION.UNIT_CODE%TYPE; a_mand FEE.FEE_INPUT_DEFINITION.MANDATORY_FLAG%TYPE; a_order NUMBER;')
    a("    k VARCHAR2(200) := p_fee_code||'#'||p_input_code;")
    a('  BEGIN')
    a('    v_inputs_checked := v_inputs_checked + 1;')
    a('    BEGIN')
    a('      SELECT i.NAME_FA,i.DATA_TYPE_CODE,i.UNIT_CODE,i.MANDATORY_FLAG,i.DISPLAY_ORDER')
    a('        INTO a_name,a_type,a_unit,a_mand,a_order')
    a('        FROM FEE_INPUT_DEFINITION i')
    a('        JOIN FEE_CALCULATION_RULE r ON r.CALCULATION_RULE_ID=i.CALCULATION_RULE_ID')
    a('        JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID')
    a('        JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID')
    a(f"       WHERE d.FEE_CODE=p_fee_code AND dv.VERSION_NO='{POLICY_VERSION_NO}' AND i.INPUT_CODE=p_input_code;")
    a('    EXCEPTION')
    a("      WHEN NO_DATA_FOUND THEN fail('INPUT',k,'ROW','present','missing'); RETURN;")
    a("      WHEN TOO_MANY_ROWS THEN fail('INPUT',k,'ROW','one row','multiple rows'); RETURN;")
    a('    END;')
    a("    check_text('INPUT',k,'NAME_FA',p_name,a_name);")
    a("    check_text('INPUT',k,'DATA_TYPE_CODE','NUMBER',a_type);")
    a("    check_text('INPUT',k,'UNIT_CODE',p_unit,a_unit);")
    a("    check_text('INPUT',k,'MANDATORY_FLAG','Y',a_mand);")
    a("    check_num('INPUT',k,'DISPLAY_ORDER',p_order,a_order);")
    a('  END;')
    a('')
    a('  PROCEDURE check_tier(p_fee_code VARCHAR2, p_tier_no NUMBER, p_name VARCHAR2, p_lower NUMBER, p_upper NUMBER, p_basis VARCHAR2, p_strategy VARCHAR2, p_fixed NUMBER, p_rate NUMBER, p_min NUMBER, p_max NUMBER) IS')
    a('    a_name FEE.FEE_CALCULATION_TIER.TIER_NAME_FA%TYPE; a_lower NUMBER; a_upper NUMBER; a_unit FEE.FEE_CALCULATION_TIER.BOUND_UNIT_CODE%TYPE; a_basis FEE.FEE_CALCULATION_TIER.TIER_BASIS_CODE%TYPE; a_strategy FEE.FEE_CALCULATION_TIER.TIER_STRATEGY_CODE%TYPE;')
    a('    a_fixed NUMBER; a_rate NUMBER; a_min NUMBER; a_max NUMBER; a_from DATE; a_to DATE;')
    a("    k VARCHAR2(200) := p_fee_code||'#TIER'||p_tier_no;")
    a('  BEGIN')
    a('    v_tiers_checked := v_tiers_checked + 1;')
    a('    BEGIN')
    a('      SELECT tr.TIER_NAME_FA,tr.LOWER_BOUND,tr.UPPER_BOUND,tr.BOUND_UNIT_CODE,tr.TIER_BASIS_CODE,tr.TIER_STRATEGY_CODE,')
    a('             tr.FIXED_AMOUNT,tr.RATE_VALUE,tr.MIN_FEE_AMOUNT,tr.MAX_FEE_AMOUNT,tr.EFFECTIVE_FROM,tr.EFFECTIVE_TO')
    a('        INTO a_name,a_lower,a_upper,a_unit,a_basis,a_strategy,a_fixed,a_rate,a_min,a_max,a_from,a_to')
    a('        FROM FEE_CALCULATION_TIER tr')
    a('        JOIN FEE_CALCULATION_RULE r ON r.CALCULATION_RULE_ID=tr.CALCULATION_RULE_ID')
    a('        JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID')
    a('        JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID')
    a(f"       WHERE d.FEE_CODE=p_fee_code AND dv.VERSION_NO='{POLICY_VERSION_NO}' AND tr.TIER_NO=p_tier_no;")
    a('    EXCEPTION')
    a("      WHEN NO_DATA_FOUND THEN fail('TIER',k,'ROW','present','missing'); RETURN;")
    a("      WHEN TOO_MANY_ROWS THEN fail('TIER',k,'ROW','one row','multiple rows'); RETURN;")
    a('    END;')
    a("    check_text('TIER',k,'TIER_NAME_FA',p_name,a_name);")
    a("    check_num('TIER',k,'LOWER_BOUND',p_lower,a_lower);")
    a("    check_num('TIER',k,'UPPER_BOUND',p_upper,a_upper);")
    a("    check_text('TIER',k,'BOUND_UNIT_CODE','IRR',a_unit);")
    a("    check_text('TIER',k,'TIER_BASIS_CODE',p_basis,a_basis);")
    a("    check_text('TIER',k,'TIER_STRATEGY_CODE',p_strategy,a_strategy);")
    a("    check_num('TIER',k,'FIXED_AMOUNT',p_fixed,a_fixed);")
    a("    check_num('TIER',k,'RATE_VALUE',p_rate,a_rate);")
    a("    check_num('TIER',k,'MIN_FEE_AMOUNT',p_min,a_min);")
    a("    check_num('TIER',k,'MAX_FEE_AMOUNT',p_max,a_max);")
    a(f"    check_date('TIER',k,'EFFECTIVE_FROM',DATE '{EFFECTIVE_FROM}',a_from);")
    a("    check_date('TIER',k,'EFFECTIVE_TO',NULL,a_to);")
    a('  END;')
    a('')
    a('  PROCEDURE check_global_counts IS')
    a('    n NUMBER;')
    a('  BEGIN')
    a(f"    SELECT COUNT(*) INTO n FROM FEE_DEFINITION WHERE CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' AND IS_ACTIVE='Y';")
    a("    check_num('GLOBAL','CBI1405','DEFINITIONS',152,n);")
    a(f"    SELECT COUNT(*) INTO n FROM FEE_CALCULATION_RULE r JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' AND r.IS_ACTIVE='Y';")
    a("    check_num('GLOBAL','CBI1405','RULES',152,n);")
    a(f"    SELECT COUNT(*) INTO n FROM FEE_RULE_COMPONENT c JOIN FEE_CALCULATION_RULE r ON r.CALCULATION_RULE_ID=c.CALCULATION_RULE_ID JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' AND c.REFERENCE_CODE LIKE 'SRC:%';")
    a("    check_num('GLOBAL','CBI1405','SOURCE_COMPONENTS',180,n);")
    a(f"    SELECT COUNT(*) INTO n FROM FEE_INPUT_DEFINITION i JOIN FEE_CALCULATION_RULE r ON r.CALCULATION_RULE_ID=i.CALCULATION_RULE_ID JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}';")
    a("    check_num('GLOBAL','CBI1405','INPUT_DEFINITIONS',66,n);")
    a(f"    SELECT COUNT(*) INTO n FROM FEE_CALCULATION_TIER tr JOIN FEE_CALCULATION_RULE r ON r.CALCULATION_RULE_ID=tr.CALCULATION_RULE_ID JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}';")
    a("    check_num('GLOBAL','CBI1405','TIERS',15,n);")
    a('  END;')
    a('')
    a('  PROCEDURE check_source_hashes IS')
    a('    a_ref VARCHAR2(500); a_hash VARCHAR2(200); a_status VARCHAR2(30);')
    a('  BEGIN')
    a(f"    SELECT DOCUMENT_REF,DOCUMENT_HASH,STATUS_CODE INTO a_ref,a_hash,a_status FROM FEE_REGULATORY_SOURCE WHERE SOURCE_CODE='{SOURCE_CODE}';")
    a(f"    IF INSTR(a_ref,'XLSX_SHA256={xlsx_hash}')=0 THEN fail('SOURCE','{SOURCE_CODE}','XLSX_SHA256','{xlsx_hash}',a_ref); END IF;")
    a(f"    check_text('SOURCE','{SOURCE_CODE}','PDF_SHA256','SHA256:{pdf_hash}',a_hash);")
    a(f"    check_text('SOURCE','{SOURCE_CODE}','STATUS_CODE','PROVISIONAL',a_status);")
    a(f"  EXCEPTION WHEN NO_DATA_FOUND THEN fail('SOURCE','{SOURCE_CODE}','ROW','present','missing');")
    a('  END;')
    a('')
    a('BEGIN')
    a("  DBMS_OUTPUT.PUT_LINE('=== CBI 1405 PROVISIONAL SOURCE -> ORACLE RECONCILIATION ===');")
    a('  check_source_hashes;')
    a('  check_global_counts;')
    a('')
    a("  DBMS_OUTPUT.PUT_LINE('--- Checking 152 tariff contracts ---');")
    for t in tariffs:
        rule_code='TARIFF_'+slug(t.tariff_code)
        a('  check_tariff('+','.join([
            sql_str(t.fee_code),sql_str(t.tariff_code),sql_str(rule_code),sql_str(t.service_name),
            sql_str(t.feature_code),sql_str(t.category_code),sql_str(t.config_hash),sql_str(t.strategy),sql_str(t.basis_type),
            sql_num(t.fixed_amount),sql_num(t.rate_value),sql_num(t.min_fee),sql_num(t.max_fee),sql_str(t.rate_period_code),str(t.component_count)
        ])+');')
    a('')
    a("  DBMS_OUTPUT.PUT_LINE('--- Checking 180 preserved source components ---');")
    for t in tariffs:
        for c in by_tariff[t.tariff_code]:
            node=component_node_type(c)
            num=component_numeric(c)
            ref=f"SRC:{c['component_id']}:{c['rule_type']}"[:190]
            desc=component_description(c)
            text=c['original_fee_text'][:490] if c['original_fee_text'] else None
            a('  check_component('+','.join([
                sql_str(t.fee_code),str(c['sequence']),sql_str(node),sql_num(num),sql_str(text),sql_str(ref),sql_str(desc)
            ])+');')
    a('')
    a("  DBMS_OUTPUT.PUT_LINE('--- Checking executable input contracts ---');")
    percentage_strategies={'PERCENTAGE','PERCENTAGE_WITH_FLOOR','PERCENTAGE_WITH_CAP','PERCENTAGE_FLOOR_CAP','ANNUALIZED_PERCENTAGE'}
    for t in tariffs:
        comps=by_tariff[t.tariff_code]
        if t.strategy in percentage_strategies:
            basis_name=comps[0]['basis'] or 'مبلغ مبنای محاسبه'
            a('  check_input('+','.join([sql_str(t.fee_code),sql_str('BASE_AMOUNT'),sql_str(basis_name),sql_str('IRR'),'1'])+');')
        ranged=[c for c in comps if c['lower_bound'] is not None or c['upper_bound'] is not None]
        if len(ranged)==3 and all(c['range_basis']=='مبلغ ارزیابی' for c in ranged):
            a('  check_input('+','.join([sql_str(t.fee_code),sql_str('APPRAISAL_AMOUNT'),sql_str('مبلغ ارزیابی'),sql_str('IRR'),'1'])+');')
            a('  check_input('+','.join([sql_str(t.fee_code),sql_str('OFFICIAL_EXPERT_TARIFF'),sql_str('تعرفه کارشناس رسمی دادگستری'),sql_str('IRR'),'2'])+');')
        elif t.strategy=='PER_UNIT':
            unit=comps[0]['unit'] or 'واحد'
            a('  check_input('+','.join([sql_str(t.fee_code),sql_str('UNIT_COUNT'),sql_str('تعداد '+unit),sql_str('COUNT'),'1'])+');')
    a('')
    a("  DBMS_OUTPUT.PUT_LINE('--- Checking 15 appraisal tiers ---');")
    for t in tariffs:
        ranged=[c for c in by_tariff[t.tariff_code] if c['lower_bound'] is not None or c['upper_bound'] is not None]
        if len(ranged)==3 and all(c['range_basis']=='مبلغ ارزیابی' for c in ranged):
            for tier_no,c in enumerate(ranged,1):
                if c['rule_type']=='FIXED_AMOUNT':
                    tier_strategy='FIXED'; fixed=c['amount']; rate=None; basis_code='WHOLE_AMOUNT'
                elif c['rule_type']=='RATE':
                    tier_strategy='COMPOSITE' if c['reference'] else 'PERCENTAGE'; fixed=None; rate=normalized_rate(c['percent'],c['permille']); basis_code='EXCESS_OVER_LOWER_BOUND'
                else:
                    tier_strategy='EXTERNAL_VALUE'; fixed=None; rate=None; basis_code='EXCESS_OVER_LOWER_BOUND'
                tier_name=(c['description'] or f'بازه {tier_no}')[:190]
                a('  check_tier('+','.join([
                    sql_str(t.fee_code),str(tier_no),sql_str(tier_name),sql_num(c['lower_bound']),sql_num(c['upper_bound']),
                    sql_str(basis_code),sql_str(tier_strategy),sql_num(fixed),sql_num(rate),sql_num(c['min_fee']),sql_num(c['max_fee'])
                ])+');')
    a('')
    a("  DBMS_OUTPUT.PUT_LINE('--- Reconciliation summary ---');")
    a("  DBMS_OUTPUT.PUT_LINE('tariffs_checked='||v_tariffs_checked||' expected=152');")
    a("  DBMS_OUTPUT.PUT_LINE('components_checked='||v_components_checked||' expected=180');")
    a("  DBMS_OUTPUT.PUT_LINE('inputs_checked='||v_inputs_checked||' expected=66');")
    a("  DBMS_OUTPUT.PUT_LINE('tiers_checked='||v_tiers_checked||' expected=15');")
    a("  DBMS_OUTPUT.PUT_LINE('mismatches='||v_errors||' expected=0');")
    a("  IF v_tariffs_checked<>152 THEN fail('GLOBAL','CBI1405','TARIFF_CALL_COUNT','152',TO_CHAR(v_tariffs_checked)); END IF;")
    a("  IF v_components_checked<>180 THEN fail('GLOBAL','CBI1405','COMPONENT_CALL_COUNT','180',TO_CHAR(v_components_checked)); END IF;")
    a("  IF v_inputs_checked<>66 THEN fail('GLOBAL','CBI1405','INPUT_CALL_COUNT','66',TO_CHAR(v_inputs_checked)); END IF;")
    a("  IF v_tiers_checked<>15 THEN fail('GLOBAL','CBI1405','TIER_CALL_COUNT','15',TO_CHAR(v_tiers_checked)); END IF;")
    a("  IF v_errors>0 THEN RAISE_APPLICATION_ERROR(-20260,'CBI 1405 source-to-Oracle reconciliation failed; mismatches='||v_errors); END IF;")
    a("  DBMS_OUTPUT.PUT_LINE('CBI Rial Fee 1405 PROVISIONAL source-to-Oracle reconciliation OK.');")
    a('END;')
    a('/')
    return '\n'.join(lines)+'\n'

def generate_verify(meta:dict[str,Any]) -> str:
    strategy_lines='\n'.join(f"PROMPT Expected {k} = {v}" for k,v in sorted(meta['rule_strategies'].items()))
    return f"""-- CBI Rial Fee 1405 PROVISIONAL verification
SET DEFINE OFF;
WHENEVER SQLERROR EXIT SQL.SQLCODE ROLLBACK;
ALTER SESSION SET CURRENT_SCHEMA = FEE;
SET SERVEROUTPUT ON;

DECLARE
  v NUMBER;
  PROCEDURE assert_eq(p_label VARCHAR2, p_actual NUMBER, p_expected NUMBER) IS
  BEGIN
    DBMS_OUTPUT.PUT_LINE(RPAD(p_label,52)||' actual='||p_actual||' expected='||p_expected);
    IF p_actual <> p_expected THEN RAISE_APPLICATION_ERROR(-20240, p_label||' mismatch'); END IF;
  END;
BEGIN
  SELECT COUNT(*) INTO v FROM FEE_REGULATORY_SOURCE WHERE SOURCE_CODE='{SOURCE_CODE}' AND STATUS_CODE='PROVISIONAL';
  assert_eq('provisional regulatory source', v, 1);

  SELECT COUNT(*) INTO v FROM FEE_FEATURE WHERE FEATURE_CODE LIKE 'CBI_RIAL_1405_%' AND IS_ACTIVE='Y';
  assert_eq('CBI1405 provisional source features', v, 9);

  SELECT COUNT(*) INTO v FROM FEE_DEFINITION WHERE CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' AND IS_ACTIVE='Y';
  assert_eq('CBI1405 provisional fee definitions', v, 152);

  SELECT COUNT(*) INTO v
    FROM FEE_DEFINITION_VERSION dv JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
    JOIN FEE_REGULATORY_SOURCE rs ON rs.REGULATORY_SOURCE_ID=dv.REGULATORY_SOURCE_ID
   WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' AND dv.VERSION_NO='{POLICY_VERSION_NO}' AND rs.SOURCE_CODE='{SOURCE_CODE}' AND dv.STATUS_CODE='ACTIVE';
  assert_eq('CBI1405 provisional definition versions', v, 152);

  SELECT COUNT(*) INTO v
    FROM FEE_CALCULATION_RULE r JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID
    JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
   WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' AND r.IS_ACTIVE='Y';
  assert_eq('CBI1405 provisional calculation rules', v, 152);

  SELECT COUNT(*) INTO v
    FROM FEE_RULE_COMPONENT c JOIN FEE_CALCULATION_RULE r ON r.CALCULATION_RULE_ID=c.CALCULATION_RULE_ID
    JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID
    JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
   WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' AND c.REFERENCE_CODE LIKE 'SRC:%';
  assert_eq('source calculation components', v, 180);

  SELECT COUNT(*) INTO v
    FROM FEE_CALCULATION_TIER tr JOIN FEE_CALCULATION_RULE r ON r.CALCULATION_RULE_ID=tr.CALCULATION_RULE_ID
    JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID
    JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
   WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}';
  assert_eq('structured appraisal tiers', v, 15);

  SELECT COUNT(*) INTO v
    FROM FEE_CALCULATION_RULE r JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID
    JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
   WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' AND r.RATE_VALUE IS NOT NULL AND (r.RATE_VALUE < 0 OR r.RATE_VALUE > 1);
  assert_eq('invalid normalized rates', v, 0);

  SELECT COUNT(*) INTO v FROM (
    SELECT FEE_CODE FROM FEE_DEFINITION WHERE CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' GROUP BY FEE_CODE HAVING COUNT(*)>1
  );
  assert_eq('duplicate provisional fee codes', v, 0);

  SELECT COUNT(*) INTO v
    FROM FEE_DEFINITION_VERSION dv JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
    JOIN FEE_FEATURE f ON f.FEE_FEATURE_ID=d.FEE_FEATURE_ID
    JOIN FEE_REGULATORY_SOURCE rs ON rs.REGULATORY_SOURCE_ID=dv.REGULATORY_SOURCE_ID
   WHERE rs.SOURCE_CODE='{OLD_SOURCE_CODE}' AND dv.VERSION_NO='{OLD_VERSION_NO}'
     AND f.FEATURE_CODE IN ({', '.join(sql_str(x) for x in OLD_RIAL_FEATURES)})
     AND dv.STATUS_CODE='SUPERSEDED' AND dv.EFFECTIVE_TO=DATE '{OLD_ARCHIVE_TO}';
  assert_eq('archived prior non-electronic rial versions', v, 156);

  SELECT COUNT(*) INTO v
    FROM FEE_DEFINITION_VERSION dv JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
    JOIN FEE_FEATURE f ON f.FEE_FEATURE_ID=d.FEE_FEATURE_ID
    JOIN FEE_REGULATORY_SOURCE rs ON rs.REGULATORY_SOURCE_ID=dv.REGULATORY_SOURCE_ID
   WHERE rs.SOURCE_CODE='{OLD_SOURCE_CODE}' AND dv.VERSION_NO='{OLD_VERSION_NO}'
     AND f.FEATURE_CODE IN ({', '.join(sql_str(x) for x in OLD_ELECTRONIC_FEATURES)})
     AND (dv.EFFECTIVE_TO IS NULL OR dv.EFFECTIVE_TO >= DATE '{EFFECTIVE_FROM}');
  assert_eq('retained prior electronic versions', v, 73);

  DBMS_OUTPUT.PUT_LINE('CBI Rial Fee 1405 PROVISIONAL verification OK.');
END;
/

{strategy_lines}
SELECT CALCULATION_STRATEGY_CODE, COUNT(*) AS RULE_COUNT
  FROM FEE_CALCULATION_RULE r
  JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID
  JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
 WHERE d.CLASSIFICATION_CODE='{CLASSIFICATION_CODE}' AND r.IS_ACTIVE='Y'
 GROUP BY CALCULATION_STRATEGY_CODE
 ORDER BY CALCULATION_STRATEGY_CODE;
"""


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--xlsx',required=True,type=Path)
    ap.add_argument('--pdf',required=True,type=Path)
    ap.add_argument('--old-clean',required=True,type=Path)
    ap.add_argument('--out',required=True,type=Path)
    args=ap.parse_args()
    args.out.mkdir(parents=True,exist_ok=True)
    xhash=hashlib.sha256(args.xlsx.read_bytes()).hexdigest(); phash=hashlib.sha256(args.pdf.read_bytes()).hexdigest()
    tariffs,components,meta=load_source(args.xlsx)

    old=list(csv.DictReader(args.old_clean.open(encoding='utf-8-sig',newline='')))
    old_rial=[r for r in old if r['feature_code'] in OLD_RIAL_FEATURES]
    old_elec=[r for r in old if r['feature_code'] in OLD_ELECTRONIC_FEATURES]
    if len(old_rial)!=156 or len(old_elec)!=73:
        raise ValueError(f'Unexpected old scope counts: rial={len(old_rial)} electronic={len(old_elec)}')

    def tariff_row(t:Tariff):
        d=asdict(t)
        for k,v in list(d.items()):
            if isinstance(v,Decimal): d[k]=str(v)
        return d
    tr=[tariff_row(t) for t in tariffs]
    write_csv(args.out/'cbi_rial_fee_1405_provisional_definitions.csv',tr,list(tr[0].keys()))
    write_csv(args.out/'cbi_rial_fee_1405_provisional_components.csv',components,list(components[0].keys()))

    archive=[]
    for r in old_rial:
        archive.append({'action':'ARCHIVE_OLD_RIAL','fee_code':r['fee_code'],'tariff_code':r['tariff_code'],'name_fa':r['name_fa'],'feature_code':r['feature_code'],'effective_to':OLD_ARCHIVE_TO})
    for r in old_elec:
        archive.append({'action':'RETAIN_OLD_ELECTRONIC','fee_code':r['fee_code'],'tariff_code':r['tariff_code'],'name_fa':r['name_fa'],'feature_code':r['feature_code'],'effective_to':''})
    for t in tariffs:
        archive.append({'action':'INSERT_NEW_PROVISIONAL','fee_code':t.fee_code,'tariff_code':t.tariff_code,'name_fa':t.service_name,'feature_code':t.feature_code,'effective_to':''})
    write_csv(args.out/'cbi_rial_fee_1405_provisional_delta.csv',archive,['action','fee_code','tariff_code','name_fa','feature_code','effective_to'])

    (args.out/'01-import-cbi-rial-fee-1405-provisional.sql').write_text(generate_sql(tariffs,components,xhash,phash),encoding='utf-8')
    (args.out/'02-verify-cbi-rial-fee-1405-provisional.sql').write_text(generate_verify(meta),encoding='utf-8')

    # FIX103: split reconciliation into a transaction-neutral standalone wrapper
    # and an installer-enforced wrapper. The former must never rollback/commit
    # unrelated caller work; the latter intentionally rolls back the import if
    # business reconciliation fails before COMMIT.
    strict_reconcile = generate_reconciliation(tariffs,components,xhash,phash)
    core_lines=[]
    for line in strict_reconcile.splitlines():
        if line in {
            'SET DEFINE OFF;',
            'WHENEVER SQLERROR EXIT SQL.SQLCODE ROLLBACK;',
            'ALTER SESSION SET CURRENT_SCHEMA = FEE;',
            'SET SERVEROUTPUT ON SIZE UNLIMITED;',
        }:
            continue
        if line == '-- This script is read-only. It raises ORA-20260 if any business-value mismatch exists.':
            core_lines.append('-- Core reconciliation contract. Wrapper controls transaction/error behavior.')
        else:
            core_lines.append(line)
    core_text='\n'.join(core_lines)+'\n'
    core_name='04-reconcile-cbi-rial-fee-1405-provisional-core.sql'
    (args.out/core_name).write_text(core_text,encoding='utf-8')

    standalone = f"""-- Standalone read-only reconciliation wrapper.
-- IMPORTANT: on mismatch this wrapper reports ORA-20260 but does NOT COMMIT or ROLLBACK
-- the caller transaction. Run it in a clean/fresh session for independent validation.
SET DEFINE OFF;
SET SERVEROUTPUT ON SIZE UNLIMITED;
WHENEVER SQLERROR CONTINUE NONE;
ALTER SESSION SET CURRENT_SCHEMA = FEE;
PROMPT === Standalone CBI Rial Fee 1405 source-to-Oracle reconciliation ===
@{core_name}
"""
    (args.out/'04-reconcile-cbi-rial-fee-1405-provisional.sql').write_text(standalone,encoding='utf-8')

    enforced = f"""-- Installer-enforced reconciliation wrapper.
-- Any mismatch must rollback the pending import transaction before SQL*Plus exits.
SET DEFINE OFF;
SET SERVEROUTPUT ON SIZE UNLIMITED;
WHENEVER SQLERROR EXIT SQL.SQLCODE ROLLBACK;
ALTER SESSION SET CURRENT_SCHEMA = FEE;
@{core_name}
"""
    (args.out/'04-reconcile-cbi-rial-fee-1405-provisional-enforced.sql').write_text(enforced,encoding='utf-8')

    diagnose = """-- Read-only installation-state diagnostic. No COMMIT/ROLLBACK is executed.
SET DEFINE OFF;
SET SERVEROUTPUT ON;
WHENEVER SQLERROR CONTINUE NONE;
ALTER SESSION SET CURRENT_SCHEMA = FEE;

PROMPT === Oracle connection context ===
SELECT SYS_CONTEXT('USERENV','DB_NAME') DB_NAME,
       SYS_CONTEXT('USERENV','CON_NAME') CON_NAME,
       SYS_CONTEXT('USERENV','SERVICE_NAME') SERVICE_NAME,
       SYS_CONTEXT('USERENV','SESSION_USER') SESSION_USER,
       SYS_CONTEXT('USERENV','CURRENT_SCHEMA') CURRENT_SCHEMA
  FROM dual;

PROMPT === CBI 1405 provisional state ===
SELECT 'SOURCE' ITEM, COUNT(*) CNT FROM FEE_REGULATORY_SOURCE WHERE SOURCE_CODE='CBI_RIAL_FEE_1405_PROVISIONAL'
UNION ALL
SELECT 'DEFINITIONS', COUNT(*) FROM FEE_DEFINITION WHERE CLASSIFICATION_CODE='CBI_1405_RIAL_PROVISIONAL'
UNION ALL
SELECT 'ACTIVE_VERSIONS', COUNT(*)
  FROM FEE_DEFINITION_VERSION dv JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
 WHERE d.CLASSIFICATION_CODE='CBI_1405_RIAL_PROVISIONAL' AND dv.VERSION_NO='1405.P1' AND dv.STATUS_CODE='ACTIVE'
UNION ALL
SELECT 'ACTIVE_RULES', COUNT(*)
  FROM FEE_CALCULATION_RULE r JOIN FEE_DEFINITION_VERSION dv ON dv.FEE_DEFINITION_VERSION_ID=r.FEE_DEFINITION_VERSION_ID
  JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
 WHERE d.CLASSIFICATION_CODE='CBI_1405_RIAL_PROVISIONAL' AND r.IS_ACTIVE='Y';

PROMPT === Prior CBI 1404 non-electronic version state ===
SELECT dv.STATUS_CODE, dv.EFFECTIVE_TO, COUNT(*) CNT
  FROM FEE_DEFINITION_VERSION dv
  JOIN FEE_DEFINITION d ON d.FEE_DEFINITION_ID=dv.FEE_DEFINITION_ID
  JOIN FEE_FEATURE f ON f.FEE_FEATURE_ID=d.FEE_FEATURE_ID
  JOIN FEE_REGULATORY_SOURCE rs ON rs.REGULATORY_SOURCE_ID=dv.REGULATORY_SOURCE_ID
 WHERE rs.SOURCE_CODE='CBI_FEE_1404_04_35500'
   AND dv.VERSION_NO='1.0'
   AND f.FEATURE_CODE IN ('CBI_GUARANTEE_FEE_GROUP','CBI_REMITTANCE_FEE_GROUP','CBI_SAFE_DEPOSIT_FEE_GROUP','CBI_SECURITIES_FEE_GROUP','CBI_BILL_COLLECTION_FEE_GROUP','CBI_CURRENT_ACCOUNT_FEE_GROUP','CBI_SAVINGS_DEPOSIT_FEE_GROUP','CBI_CERTIFICATE_FEE_GROUP','CBI_STATEMENT_FEE_GROUP','CBI_ACCOUNT_SERVICE_FEE_GROUP','CBI_APPRAISAL_FEE_GROUP','CBI_CREDIT_FEE_GROUP','CBI_OTHER_SERVICE_FEE_GROUP')
 GROUP BY dv.STATUS_CODE,dv.EFFECTIVE_TO
 ORDER BY dv.STATUS_CODE,dv.EFFECTIVE_TO;
"""
    (args.out/'05-diagnose-cbi-rial-fee-1405-state.sql').write_text(diagnose,encoding='utf-8')

    (args.out/'00-install-cbi-rial-fee-1405-provisional.sql').write_text("""SET ECHO ON\nSET SERVEROUTPUT ON SIZE UNLIMITED\nWHENEVER SQLERROR EXIT SQL.SQLCODE ROLLBACK;\nPROMPT === Importing CBI Rial Fee 1405 PROVISIONAL ===\n@01-import-cbi-rial-fee-1405-provisional.sql\nPROMPT === Structural verification ===\n@02-verify-cbi-rial-fee-1405-provisional.sql\nPROMPT === Source-to-Oracle business-value reconciliation ===\n@04-reconcile-cbi-rial-fee-1405-provisional-enforced.sql\nCOMMIT;\nPROMPT === CBI Rial Fee 1405 PROVISIONAL committed successfully ===\n""",encoding='utf-8')
    (args.out/'03-finalize-official-reference-template.sql').write_text("""-- TEMPLATE ONLY. Fill values from the official CBI cover letter before execution.\n-- This intentionally does not modify tariff calculation data.\nSET DEFINE ON;\nDEFINE OFFICIAL_SOURCE_CODE = 'CBI_RIAL_FEE_1405_OFFICIAL';\nDEFINE OFFICIAL_CIRCULAR_NO = 'REPLACE_ME';\nDEFINE OFFICIAL_ISSUE_DATE = 'YYYY-MM-DD';\nDEFINE OFFICIAL_EFFECTIVE_FROM = 'YYYY-MM-DD';\nPROMPT Review and replace the provisional regulatory source/policy metadata manually after official verification.\n""",encoding='utf-8')

    manifest={
        'classification_code':CLASSIFICATION_CODE,'source_code':SOURCE_CODE,'source_status':'PROVISIONAL',
        'circular_no':CIRCULAR_NO,'provisional_effective_from':EFFECTIVE_FROM,'provisional_archive_to':OLD_ARCHIVE_TO,
        'xlsx_sha256':xhash,'pdf_sha256':phash,'new_tariffs':len(tariffs),'source_components':len(components),
        'old_rial_versions_to_archive':len(old_rial),'old_electronic_versions_retained':len(old_elec),
        'strategy_counts':meta['rule_strategies'],'section_counts':meta['sections'],
        'note':'Official circular number/date were not supplied; metadata is deliberately provisional.'
    }
    (args.out/'cbi_rial_fee_1405_provisional_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    readme=f"""# CBI Rial Banking Fees 1405 — Provisional Versioned Import\n\nSource workbook SHA256: `{xhash}`  \nSource PDF SHA256: `{phash}`\n\n## Scope\n- 152 new rial-banking tariff definitions from the supplied 8-page attachment.\n- {len(components)} structured source calculation components.\n- 156 prior **non-electronic** CBI-1404 fee definition versions are closed on {OLD_ARCHIVE_TO} and marked `SUPERSEDED`.\n- 73 prior **electronic** CBI-1404 definitions are explicitly retained.\n- No physical DELETE is executed. Historical definitions/rules remain queryable by effective date and regulatory source.\n\n## Provisional regulatory metadata\nThe official cover letter/circular number and effective date were not supplied. For prototype execution only, this package uses:\n- `SOURCE_CODE={SOURCE_CODE}`\n- `CIRCULAR_NO={CIRCULAR_NO}`\n- `STATUS_CODE=PROVISIONAL`\n- `EFFECTIVE_FROM={EFFECTIVE_FROM}` (derived from the supplied workbook date, **not claimed as the official effective date**)\n\nReplace/finalize these values after receiving the official CBI letter.\n\n## Calculation policy\n- Simple fixed/rate/per-unit rows are mapped to executable `FEE_CALCULATION_RULE` fields.\n- Percent values are normalized from human percent to decimal (`0.5% -> 0.005`).\n- Composite/conditional/reference/formula rows are not guessed. They use `COMPOSITE` or `EXTERNAL_VALUE`, and every source component is preserved in `FEE_RULE_COMPONENT` with source condition/reference text.\n- Source component count is preserved exactly.\n\n## Install\nRun `00-install-cbi-rial-fee-1405-provisional.sql`. It imports, performs structural verification, runs the full source-to-Oracle business-value reconciliation, then commits. Any SQL error or reconciliation mismatch rolls back before commit. For an already-installed database, first run `05-diagnose-cbi-rial-fee-1405-state.sql` if installation state is uncertain, then run `04-reconcile-cbi-rial-fee-1405-provisional.sql` directly. The standalone reconciliation wrapper is transaction-neutral: it never commits or rolls back caller work.\n"""
    (args.out/'README.md').write_text(readme,encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
